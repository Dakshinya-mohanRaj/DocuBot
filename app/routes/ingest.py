import csv
import io
import re
import uuid
from pathlib import PurePosixPath, PureWindowsPath
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field

from app.core.chunking import chunk_text
from app.core.dependencies import CurrentUser, require_csrf
from app.core.vector_store import (
    add_chunks,
    collection_count,
    delete_document,
    document_exists,
    list_documents,
)

# Optional heavy dependencies — only required for their respective file types.
# Imported at module level so linters can resolve them; missing installs are
# caught at the point of use with a clear 400 error message.
_pypdf: Any = None
try:
    import pypdf as _pypdf
except ImportError:
    pass

_docx: Any = None
try:
    import docx as _docx
except ImportError:
    pass

_openpyxl: Any = None
try:
    import openpyxl as _openpyxl
except ImportError:
    pass


router = APIRouter(dependencies=[Depends(require_csrf)])

MAX_UPLOAD_BYTES = 10 * 1024 * 1024  # 10 MB

TEXT_EXTENSIONS = (".txt", ".md", ".json", ".log")
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".xlsx", ".csv") + TEXT_EXTENSIONS

_UNSAFE_CHARS = re.compile(r"[\x00-\x1f<>:\"|?*\\/]+")


def sanitize_doc_id(filename: str) -> str:
    """Derive a safe document id from a user-supplied filename.

    Strips any directory component and control/path characters so a crafted
    filename cannot escape its namespace or collide oddly. The original name is
    still stored separately as display metadata.
    """
    name = PureWindowsPath(filename).name
    name = PurePosixPath(name).name
    name = _UNSAFE_CHARS.sub("_", name).strip("._ ")
    if not name:
        name = "document"
    return name[:200]


class IngestTextRequest(BaseModel):
    doc_id: str | None = Field(default=None, max_length=200)
    text: str = Field(min_length=1, max_length=2_000_000)


@router.post("/ingest/text")
def ingest_text(payload: IngestTextRequest, user: CurrentUser) -> dict:
    if not payload.text.strip():
        raise HTTPException(status_code=400, detail="text cannot be empty")

    doc_id = sanitize_doc_id(payload.doc_id) if payload.doc_id else str(uuid.uuid4())[:8]
    chunks = chunk_text(payload.text)
    if not chunks:
        raise HTTPException(status_code=400, detail="text produced no usable content")

    replaced = document_exists(doc_id, user_id=user.id)
    added = add_chunks(
        doc_id, chunks, filename=payload.doc_id or doc_id, user_id=user.id
    )

    return {
        "doc_id": doc_id,
        "chunks_added": added,
        "replaced": replaced,
        "total_chunks_in_store": collection_count(user.id),
    }


@router.get("/ingest/documents")
def ingest_documents(
    user: CurrentUser,
    recent: bool = Query(default=False, description="Sort newest-first by upload time."),
) -> dict:
    """List every document the authenticated user has uploaded."""
    documents = list_documents(user_id=user.id)
    if recent:
        documents.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    return {
        "documents": documents,
        "document_count": len(documents),
        "total_chunks_in_store": collection_count(user.id),
    }


def parse_file_blocks(filename: str, raw: bytes) -> list[tuple[str, int | None]]:
    """Parse a file into (text, page) blocks.

    ``page`` is the 1-based PDF page number, or ``None`` for formats without
    pages. Chunking each block separately keeps citations page-accurate.
    """
    lower = filename.lower()

    # 1. PDF Documents (.pdf)
    if lower.endswith(".pdf"):
        if _pypdf is None:
            raise HTTPException(status_code=400, detail="PDF parsing requires 'pypdf'.")
        try:
            reader = _pypdf.PdfReader(io.BytesIO(raw))
            blocks: list[tuple[str, int | None]] = []
            for page_number, page in enumerate(reader.pages, start=1):
                try:
                    text = page.extract_text() or ""
                except Exception:
                    text = ""
                blocks.append((text, page_number))
        except Exception:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read PDF '{filename}'. It may be corrupted or password-protected.",
            )
        if not any(text.strip() for text, _ in blocks):
            raise HTTPException(
                status_code=400,
                detail=(
                    f"'{filename}' contains no extractable text. It looks like a "
                    "scanned or image-only PDF, which OCR is needed to read."
                ),
            )
        return blocks

    # 2. Word Documents (.docx)
    if lower.endswith(".docx"):
        if _docx is None:
            raise HTTPException(status_code=400, detail="Word parsing requires 'python-docx'.")
        try:
            document = _docx.Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
            return [(text, None)]
        except Exception:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read Word file '{filename}'.",
            )

    # 3. Excel Spreadsheets (.xlsx)
    if lower.endswith(".xlsx"):
        if _openpyxl is None:
            raise HTTPException(status_code=400, detail="Excel parsing requires 'openpyxl'.")
        try:
            wb = _openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
            lines: list[str] = []
            for sheet in wb.worksheets:
                lines.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    row_str = [str(cell) if cell is not None else "" for cell in row]
                    if any(row_str):
                        lines.append(" | ".join(row_str))
            return [("\n".join(lines), None)]
        except Exception:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read Excel file '{filename}'.",
            )

    # 4. CSV Files (.csv)
    if lower.endswith(".csv"):
        try:
            decoded = raw.decode("utf-8", errors="ignore")
            reader = csv.reader(io.StringIO(decoded))
            text = "\n".join(" | ".join(row) for row in reader if any(row))
            return [(text, None)]
        except Exception:
            raise HTTPException(
                status_code=400,
                detail=f"Could not read CSV file '{filename}'.",
            )

    # 5. Text & Markdown (.txt, .md, .json, .log)
    if lower.endswith(TEXT_EXTENSIONS):
        return [(raw.decode("utf-8", errors="ignore"), None)]

    raise HTTPException(
        status_code=400,
        detail=f"Unsupported file format '{filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}",
    )


def parse_file_content(filename: str, raw: bytes) -> str:
    """Backwards-compatible helper: flattened text for a single file."""
    return "\n\n".join(text for text, _ in parse_file_blocks(filename, raw))


@router.post("/ingest/file")
async def ingest_file(user: CurrentUser, file: UploadFile = File(...)) -> dict:
    filename = (file.filename or "").strip()
    if not filename:
        raise HTTPException(status_code=400, detail="Uploaded file has no name.")

    if not filename.lower().endswith(SUPPORTED_EXTENSIONS):
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{filename}'. Supported: {', '.join(SUPPORTED_EXTENSIONS)}",
        )

    # Read at most one byte over the limit so an oversized upload is rejected
    # without buffering the whole body into memory.
    raw = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File too large. Maximum allowed size is 10 MB.",
        )
    if not raw:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    try:
        blocks = parse_file_blocks(filename, raw)
    except HTTPException:
        raise

    chunks: list[str] = []
    chunk_meta: list[dict] = []
    for text, page in blocks:
        for chunk in chunk_text(text):
            chunks.append(chunk)
            chunk_meta.append({"page": page} if page is not None else {})

    if not chunks:
        raise HTTPException(
            status_code=400, detail="Uploaded file contained no usable text."
        )

    # The stored doc_id is a sanitized name; the display name is preserved.
    doc_id = sanitize_doc_id(filename)
    replaced = document_exists(doc_id, user_id=user.id)
    added = add_chunks(
        doc_id, chunks, filename=filename, metadatas=chunk_meta, user_id=user.id
    )

    return {
        "doc_id": doc_id,
        "chunks_added": added,
        "replaced": replaced,
        "total_chunks_in_store": collection_count(user.id),
    }


@router.delete("/ingest/{doc_id:path}")
def delete_ingested_doc(doc_id: str, user: CurrentUser) -> dict:
    deleted_count = delete_document(doc_id, user_id=user.id)
    if deleted_count == 0:
        raise HTTPException(status_code=404, detail=f"No document named '{doc_id}' was found.")
    return {
        "message": f"Successfully deleted document '{doc_id}'",
        "chunks_deleted": deleted_count,
        "total_chunks_in_store": collection_count(user.id),
    }
